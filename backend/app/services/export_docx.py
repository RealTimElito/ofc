"""Convert Markdown report body to a .docx Word document."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Optional

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.services.theme import apply_theme_to_document, empty_theme, theme_has_visuals


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_UL_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
_OL_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_HR_RE = re.compile(r"^(\*{3,}|-{3,}|_{3,})\s*$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
_INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\*[^*]+\*|_[^_]+_|\[[^\]]+\]\([^)]+\))"
)
_LEADING_H1_RE = re.compile(r"^#\s+(.+?)\s*$")


def _norm_heading(text: str) -> str:
    return " ".join((text or "").split()).casefold()


def strip_redundant_title_heading(title: str, body_md: str) -> str:
    """Drop a leading ATX H1 that duplicates the report title (avoids double H1)."""
    if not body_md:
        return body_md or ""
    lines = body_md.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i >= len(lines):
        return body_md
    match = _LEADING_H1_RE.match(lines[i].strip())
    if not match:
        return body_md
    if _norm_heading(match.group(1)) != _norm_heading(title or ""):
        return body_md
    del lines[i]
    if i < len(lines) and not lines[i].strip():
        del lines[i]
    return "\n".join(lines)


def _set_run_font(
    run,
    *,
    mono: bool = False,
    body_font: str = "Calibri",
    size_pt: float | None = None,
    set_size: bool = True,
) -> None:
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    elif set_size:
        run.font.size = Pt(10 if mono else 11)
    name = "Courier New" if mono else body_font
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def _add_inline_runs(
    paragraph, text: str, *, body_font: str = "Calibri", size_pt: float | None = None
) -> None:
    if not text:
        return
    pos = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _set_run_font(run, body_font=body_font, size_pt=size_pt)
        token = match.group(0)
        if token.startswith("**") or token.startswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_run_font(run, body_font=body_font, size_pt=size_pt)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, mono=True)
        elif token.startswith("*") or token.startswith("_"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            _set_run_font(run, body_font=body_font, size_pt=size_pt)
        elif token.startswith("["):
            label, _, rest = token[1:].partition("](")
            url = rest[:-1] if rest.endswith(")") else rest
            run = paragraph.add_run(label or url)
            run.font.color.rgb = RGBColor(0x0F, 0x43, 0x38)
            run.underline = True
            _set_run_font(run, body_font=body_font, size_pt=size_pt)
            if url:
                paragraph.add_run(f" ({url})")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, body_font=body_font, size_pt=size_pt)


def _split_table_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [c.strip() for c in raw.split("|")]


def _add_table(
    doc: Document,
    header: list[str],
    rows: list[list[str]],
    *,
    body_font: str,
    size_pt: float | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.bold = True
        _set_run_font(run, body_font=body_font, size_pt=size_pt)
    for r_idx, row in enumerate(rows):
        for c_idx in range(len(header)):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            value = row[c_idx] if c_idx < len(row) else ""
            _add_inline_runs(p, value, body_font=body_font, size_pt=size_pt)
    doc.add_paragraph("")


def _add_code_block(doc: Document, lines: Iterable[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_run_font(run, mono=True)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph("")


def markdown_to_docx_bytes(
    title: str,
    body_md: str,
    *,
    theme: Optional[dict[str, Any]] = None,
    theme_assets_dir: Optional[Path] = None,
) -> bytes:
    """Build a .docx from Markdown. Legacy .doc is not supported."""
    theme = theme or empty_theme()
    body_md = strip_redundant_title_heading(title or "", body_md or "")
    body_font = theme.get("body_font") or "Calibri"
    heading_font = theme.get("heading_font") or body_font
    body_size = theme.get("body_size_pt") or 11
    title_size = theme.get("title_size_pt") or theme.get("heading_size_pt") or 18
    heading_size = theme.get("heading_size_pt")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = body_font
    style.font.size = Pt(float(body_size))

    if theme_has_visuals(theme):
        apply_theme_to_document(doc, theme, theme_assets_dir or Path("."))

    heading = doc.add_heading(title or "Report", level=0)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    for run in heading.runs:
        _set_run_font(run, body_font=heading_font, size_pt=float(title_size))

    lines = (body_md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    para_buf: list[str] = []

    def flush_para() -> None:
        nonlocal para_buf
        if not para_buf:
            return
        text = " ".join(para_buf).strip()
        para_buf = []
        if not text:
            return
        p = doc.add_paragraph()
        _add_inline_runs(p, text, body_font=body_font, size_pt=float(body_size))

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_para()
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            _add_code_block(doc, code_lines)
            if i < len(lines):
                i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_SEP_RE.match(
            lines[i + 1].strip()
        ):
            flush_para()
            header = _split_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            if header:
                _add_table(
                    doc, header, rows, body_font=body_font, size_pt=float(body_size)
                )
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        if _HR_RE.match(stripped):
            flush_para()
            p = doc.add_paragraph("―" * 24)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        heading_match = _HEADING_RE.match(stripped)
        if heading_match:
            flush_para()
            level = min(len(heading_match.group(1)), 4)
            h = doc.add_heading(heading_match.group(2).strip(), level=level)
            size = float(heading_size) if heading_size else None
            for run in h.runs:
                _set_run_font(
                    run,
                    body_font=heading_font,
                    size_pt=size,
                    set_size=size is not None,
                )
            i += 1
            continue

        ul = _UL_RE.match(line)
        if ul:
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(
                p, ul.group(3), body_font=body_font, size_pt=float(body_size)
            )
            i += 1
            continue

        ol = _OL_RE.match(line)
        if ol:
            flush_para()
            p = doc.add_paragraph(style="List Number")
            _add_inline_runs(
                p, ol.group(3), body_font=body_font, size_pt=float(body_size)
            )
            i += 1
            continue

        para_buf.append(stripped)
        i += 1

    flush_para()

    if len(doc.paragraphs) <= 1 and not (body_md or "").strip():
        doc.add_paragraph("(empty report)")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
