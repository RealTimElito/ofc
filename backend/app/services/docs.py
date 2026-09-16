"""Extract readable text from uploaded documents."""

from __future__ import annotations

from pathlib import Path

MAX_FILE_CHARS = 80_000

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".rst",
}


def _truncate(text: str) -> str:
    if len(text) > MAX_FILE_CHARS:
        return text[:MAX_FILE_CHARS] + f"\n\n[Truncated; original length {len(text)} chars]"
    return text


def extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _truncate("\n\n".join(parts))


def read_upload_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        raw = path.read_text(encoding="utf-8", errors="replace")
        return _truncate(raw)
    if suffix == ".docx":
        try:
            return extract_docx(path)
        except Exception as exc:  # noqa: BLE001
            return f"[Could not read Word document {path.name}: {exc}]"
    if suffix == ".doc":
        return (
            f"[Legacy .doc is not supported inline. Convert to .docx or .md: {path.name}]"
        )
    return f"[Binary or unsupported file skipped for inline context: {path.name}]"
