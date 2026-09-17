"""Extract and apply visual themes from Word .docx examples."""

from __future__ import annotations

import json
import mimetypes
import shutil
from pathlib import Path
from typing import Any, Optional

from docx import Document as DocxDocument
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

THEME_KEYS = (
    "heading_font",
    "body_font",
    "header_text",
    "footer_text",
    "header_logo",
    "footer_logo",
    "source_label",
    "margin_top_in",
    "margin_bottom_in",
    "margin_left_in",
    "margin_right_in",
    "body_size_pt",
    "heading_size_pt",
    "title_size_pt",
)

_EMU_PER_INCH = 914400.0


def empty_theme() -> dict[str, Any]:
    return {
        "heading_font": None,
        "body_font": None,
        "header_text": "",
        "footer_text": "",
        "header_logo": None,
        "footer_logo": None,
        "source_label": "",
        "margin_top_in": None,
        "margin_bottom_in": None,
        "margin_left_in": None,
        "margin_right_in": None,
        "body_size_pt": None,
        "heading_size_pt": None,
        "title_size_pt": None,
    }


def theme_has_visuals(theme: dict[str, Any]) -> bool:
    """True when export should apply theme chrome/fonts/margins."""
    if theme.get("body_font") or theme.get("heading_font"):
        return True
    if theme.get("header_text") or theme.get("footer_text"):
        return True
    if theme.get("header_logo") or theme.get("footer_logo"):
        return True
    for key in (
        "margin_top_in",
        "margin_bottom_in",
        "margin_left_in",
        "margin_right_in",
        "body_size_pt",
        "heading_size_pt",
        "title_size_pt",
    ):
        if theme.get(key) is not None:
            return True
    return False


def parse_theme_json(raw: Optional[str]) -> dict[str, Any]:
    theme = empty_theme()
    if not raw:
        return theme
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return theme
    if not isinstance(data, dict):
        return theme
    for key in THEME_KEYS:
        if key in data and data[key] is not None:
            theme[key] = data[key]
    return theme


def theme_to_json(theme: dict[str, Any]) -> str:
    return json.dumps({k: theme.get(k) for k in THEME_KEYS}, ensure_ascii=False)


def _font_name_from_style(style) -> Optional[str]:
    try:
        name = style.font.name
        if name:
            return str(name)
    except Exception:  # noqa: BLE001
        pass
    try:
        r_pr = style._element.rPr  # type: ignore[attr-defined]
        if r_pr is not None and r_pr.rFonts is not None:
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                val = r_pr.rFonts.get(qn(f"w:{attr}"))
                if val:
                    return str(val)
    except Exception:  # noqa: BLE001
        pass
    return None


def _pt_from_style(style) -> Optional[float]:
    try:
        size = style.font.size
        if size is not None:
            return round(float(size.pt), 2)
    except Exception:  # noqa: BLE001
        pass
    return None


def _emu_to_inches(emu) -> Optional[float]:
    if emu is None:
        return None
    try:
        return round(int(emu) / _EMU_PER_INCH, 4)
    except (TypeError, ValueError):
        return None


def _collect_paragraph_text(paragraphs) -> str:
    parts: list[str] = []
    for para in paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _image_ext(content_type: str, blob: bytes) -> str:
    ext = mimetypes.guess_extension(content_type or "") or ""
    if ext in {".jpe", ".jpeg"}:
        ext = ".jpg"
    if not ext:
        if blob.startswith(b"\x89PNG"):
            ext = ".png"
        elif blob[:3] == b"\xff\xd8\xff":
            ext = ".jpg"
        elif blob[:6] in (b"GIF87a", b"GIF89a"):
            ext = ".gif"
        else:
            ext = ".bin"
    return ext


def _save_first_image(part, dest: Path) -> Optional[str]:
    """Save the first image from a header/footer part; return relative filename."""
    for rel in part.rels.values():
        reltype = getattr(rel, "reltype", "") or ""
        if "image" not in reltype:
            continue
        try:
            blob = rel.target_part.blob
            content_type = getattr(rel.target_part, "content_type", "") or ""
        except Exception:  # noqa: BLE001
            continue
        if not blob:
            continue
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest = dest.with_suffix(_image_ext(content_type, blob))
        dest.write_bytes(blob)
        return dest.name
    return None


def extract_theme_from_docx(
    docx_path: Path,
    assets_dir: Path,
    *,
    source_label: str = "",
) -> dict[str, Any]:
    """Read fonts, header/footer text, logos, margins, and sizes from a .docx."""
    if assets_dir.exists():
        shutil.rmtree(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)

    doc = DocxDocument(str(docx_path))
    theme = empty_theme()
    theme["source_label"] = source_label or docx_path.name

    theme["body_font"] = _font_name_from_style(doc.styles["Normal"])
    theme["body_size_pt"] = _pt_from_style(doc.styles["Normal"])
    for style_name in ("Heading 1", "Title", "Heading 2"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        font = _font_name_from_style(style)
        if font and not theme["heading_font"]:
            theme["heading_font"] = font
        size = _pt_from_style(style)
        if size is not None:
            if style_name == "Title" and theme["title_size_pt"] is None:
                theme["title_size_pt"] = size
            elif style_name.startswith("Heading") and theme["heading_size_pt"] is None:
                theme["heading_size_pt"] = size
    if not theme["heading_font"]:
        theme["heading_font"] = theme["body_font"]
    if theme["title_size_pt"] is None and theme["heading_size_pt"] is not None:
        theme["title_size_pt"] = theme["heading_size_pt"]

    # Fallback: first body run fonts
    if not theme["body_font"]:
        for para in doc.paragraphs[:20]:
            for run in para.runs:
                if run.font.name:
                    theme["body_font"] = run.font.name
                    break
            if theme["body_font"]:
                break

    header_texts: list[str] = []
    footer_texts: list[str] = []
    header_logo: Optional[str] = None
    footer_logo: Optional[str] = None

    for section in doc.sections:
        if theme["margin_top_in"] is None:
            theme["margin_top_in"] = _emu_to_inches(section.top_margin)
            theme["margin_bottom_in"] = _emu_to_inches(section.bottom_margin)
            theme["margin_left_in"] = _emu_to_inches(section.left_margin)
            theme["margin_right_in"] = _emu_to_inches(section.right_margin)
        header = section.header
        footer = section.footer
        ht = _collect_paragraph_text(header.paragraphs)
        ft = _collect_paragraph_text(footer.paragraphs)
        if ht:
            header_texts.append(ht)
        if ft:
            footer_texts.append(ft)
        if header_logo is None:
            header_logo = _save_first_image(
                header.part, assets_dir / "header_logo"
            )
        if footer_logo is None:
            footer_logo = _save_first_image(
                footer.part, assets_dir / "footer_logo"
            )

    theme["header_text"] = header_texts[0] if header_texts else ""
    theme["footer_text"] = footer_texts[0] if footer_texts else ""
    theme["header_logo"] = header_logo
    theme["footer_logo"] = footer_logo
    return theme


def _apply_font_to_run(run, font_name: Optional[str], size_pt: float | None = None) -> None:
    if font_name:
        run.font.name = font_name
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:eastAsia"), font_name)
        r_fonts.set(qn("w:cs"), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _apply_font_to_style(style, font_name: Optional[str], size_pt: float | None = None) -> None:
    if not font_name and size_pt is None:
        return
    try:
        if font_name:
            style.font.name = font_name
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.get_or_add_rFonts()
            r_fonts.set(qn("w:ascii"), font_name)
            r_fonts.set(qn("w:hAnsi"), font_name)
            r_fonts.set(qn("w:eastAsia"), font_name)
            r_fonts.set(qn("w:cs"), font_name)
        if size_pt is not None:
            style.font.size = Pt(size_pt)
    except Exception:  # noqa: BLE001
        pass


def apply_theme_to_document(
    doc: DocxDocument,
    theme: dict[str, Any],
    assets_dir: Path,
) -> None:
    """Apply fonts, margins, and header/footer chrome to an open Document."""
    body_font = theme.get("body_font") or "Calibri"
    heading_font = theme.get("heading_font") or body_font
    body_size = theme.get("body_size_pt")
    heading_size = theme.get("heading_size_pt")
    title_size = theme.get("title_size_pt") or heading_size

    _apply_font_to_style(
        doc.styles["Normal"],
        body_font,
        float(body_size) if body_size is not None else 11,
    )
    _apply_font_to_style(
        doc.styles["Title"],
        heading_font,
        float(title_size) if title_size is not None else None,
    )
    for level in range(1, 5):
        try:
            style = doc.styles[f"Heading {level}"]
        except KeyError:
            continue
        size = float(heading_size) if heading_size is not None else None
        _apply_font_to_style(style, heading_font, size)

    if not doc.sections:
        return
    section = doc.sections[0]
    for attr, key in (
        ("top_margin", "margin_top_in"),
        ("bottom_margin", "margin_bottom_in"),
        ("left_margin", "margin_left_in"),
        ("right_margin", "margin_right_in"),
    ):
        inches = theme.get(key)
        if inches is None:
            continue
        try:
            setattr(section, attr, Inches(float(inches)))
        except (TypeError, ValueError):
            continue

    header = section.header
    footer = section.footer

    def _fill_chrome(
        container,
        text: str,
        logo_name: Optional[str],
        *,
        align=WD_PARAGRAPH_ALIGNMENT.LEFT,
    ) -> None:
        while len(container.paragraphs) > 1:
            p = container.paragraphs[-1]
            p._element.getparent().remove(p._element)
        para = container.paragraphs[0]
        para.clear()
        para.alignment = align
        logo_path = assets_dir / logo_name if logo_name else None
        if logo_path and logo_path.is_file():
            run = para.add_run()
            try:
                run.add_picture(str(logo_path), width=Inches(1.15))
            except Exception:  # noqa: BLE001
                pass
            if text:
                # Keep logo + caption on one line when practical
                run = para.add_run(f"  {text}")
                _apply_font_to_run(run, body_font, 9)
        elif text:
            run = para.add_run(text)
            _apply_font_to_run(run, body_font, 9)

    _fill_chrome(
        header,
        theme.get("header_text") or "",
        theme.get("header_logo"),
        align=WD_PARAGRAPH_ALIGNMENT.LEFT,
    )
    _fill_chrome(
        footer,
        theme.get("footer_text") or "",
        theme.get("footer_logo"),
        align=WD_PARAGRAPH_ALIGNMENT.CENTER,
    )
